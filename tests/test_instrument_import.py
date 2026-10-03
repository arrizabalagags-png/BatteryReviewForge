"""Synthetic compatibility and data-preservation gates; no author data fixtures."""
import csv
import importlib.metadata
import importlib.util
import io
import json
import os
from pathlib import Path
import struct
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

SCRIPT = Path(__file__).resolve().parents[1] / 'skills/voltpeer-data/scripts/instrument_import.py'
spec = importlib.util.spec_from_file_location('instrument_import', SCRIPT)
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
native_spec = importlib.util.spec_from_file_location('neware_ndax14', SCRIPT.with_name('neware_ndax14.py'))
native = importlib.util.module_from_spec(native_spec)
native_spec.loader.exec_module(native)


def has(name):
    return importlib.util.find_spec(name) is not None


def has_pinned_reader():
    try:
        return importlib.metadata.version('NewareNDA') == native.READER_VERSION
    except importlib.metadata.PackageNotFoundError:
        return False


def synthetic_ndax(path, zero_voltage=True, missing_run=False, version=14, auxiliary=False, unreferenced_zero=False):
    """Build a minimal layout from the declared structs, with known quantities.

    Includes a repeated Index=1 timestamp and a charge/discharge transition;
    neither may be merged into another observation by preparation.
    """
    def block(file_type):
        data = bytearray(8192)
        data[0], data[2] = file_type, version
        return data
    measurements, run, steps = block(1), block(18), block(7)
    for ordinal, (voltage, current) in enumerate([(3.0, .001), (0.0 if zero_voltage else 3.1, .001), (3.2, -.001)], 1):
        struct.pack_into('<ff', measurements, 4096 + 132 + (ordinal - 1) * 8, voltage, current)
    run_values = [
        (0, 0, 0, 0, 0, 1000, 1700000000, 1, 1, 0, b'\x00' * 8),
        (1000, .000001, 0, .000003, 0, 1000, 1700000001, 1, 1, 0, b'\x00' * 8),
        (2000, .000002, 0, .000006, 0, 1000, 1700000002, 1, 2, 0, b'\x00' * 8),
        (1000, 0, .000001, 0, .0000032, 1000, 1700000003, 2, 3, 0, b'\x00' * 8),
    ]
    if missing_run:
        run_values.pop()
    if unreferenced_zero:
        run_values = [values for values in run_values if values[8] != 2]
    for ordinal, values in enumerate(run_values):
        struct.pack_into('<ixffff8xiiiih8s', run, 4096 + 132 + ordinal * struct.calcsize('<ixffff8xiiiih8s'), *values)
    for ordinal, values in enumerate([(0, 1, b'\x00' * 16, 1, b'\x00' * 12), (0, 2, b'\x00' * 16, 2, b'\x00' * 12)]):
        struct.pack_into('<ii16sb12s', steps, 4096 + 132 + ordinal * 37, *values)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('data.ndc', measurements)
        archive.writestr('data_runInfo.ndc', run)
        archive.writestr('data_step.ndc', steps)
        if auxiliary:
            archive.writestr('data_AUX_1_1_1.ndc', block(5))


class InstrumentImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def make(self, text, extension='.csv', encoding='utf-8'):
        path = self.folder / ('input' + extension)
        path.write_text(text, encoding=encoding)
        return path

    def output_rows(self, result):
        with (self.folder / 'prepared' / result['tables'][0]['output']).open(encoding='utf-8', newline='') as stream:
            return list(csv.DictReader(stream))

    def test_chinese_gb18030_raw_units_and_source_bytes(self):
        source = self.make('循环号,放电比容量(mAh/g),库伦效率(%)\n1,180.10,100.5\n2,0,98.2\n', encoding='gb18030')
        before = source.read_bytes()
        result = module.prepare(source, self.folder / 'prepared')
        rows = self.output_rows(result)
        self.assertEqual(source.read_bytes(), before)
        self.assertEqual((self.folder / 'prepared/source-original.csv').read_bytes(), before)
        self.assertEqual(rows[0]['放电比容量(mAh/g)'], '180.10')
        self.assertEqual(rows[0]['vp_discharge_capacity_mah_g'], '180.10')
        self.assertEqual(rows[0]['vp_ce_pct'], '100.5')
        self.assertEqual(rows[1]['vp_discharge_capacity_mah_g'], '0')
        self.assertTrue(result['tables'][0]['numeric_ready'])
        self.assertFalse(result['ready_for_plot'])
        self.assertEqual(result['tables'][0]['role'], 'cycle')

    def test_tsv_utf16_and_duration(self):
        source = self.make('Record\tTime(h:min:s)\tVoltage(mV)\tCurrent(µA)\n1\t25:00:00.250\t3500\t-125\n', '.txt', 'utf-16')
        result = module.prepare(source, self.folder / 'prepared')
        row = self.output_rows(result)[0]
        self.assertEqual(row['vp_time_s'], '90000.250')
        self.assertEqual(row['vp_voltage_v'], '3.500')
        self.assertEqual(row['vp_current_a'], '-0.000125')
        self.assertEqual(row['Current(µA)'], '-125')

    def test_semicolon_quoted_values_and_utf8_bom(self):
        source = self.make('Cycle;Capacity(Ah);Comment\n1;0.12;"keep; this"\n', encoding='utf-8-sig')
        result = module.prepare(source, self.folder / 'prepared')
        row = self.output_rows(result)[0]
        self.assertEqual(row['Comment'], 'keep; this')
        self.assertEqual(row['vp_capacity_mah'], '120.00')

    def test_multiple_headers_remain_independent_and_metadata_is_copied(self):
        source = self.make('Operator,synthetic\nCycle,Capacity(mAh)\n1,2\n\nCycle,Capacity(mAh)\n2,3\n')
        result = module.prepare(source, self.folder / 'prepared')
        self.assertEqual(len(result['tables']), 2)
        self.assertEqual([t['source_rows'] for t in result['tables']], [[3], [6]])
        self.assertEqual(result['tables'][0]['blank_source_rows'], [4])
        self.assertIn('Operator,synthetic', (self.folder / 'prepared/source-original.csv').read_text(encoding='utf-8'))

    def test_repeated_alias_units_block_inference(self):
        source = self.make('Record,Time(s),Voltage(V),电压(mV)\n1,0,3.6,3600\n')
        result = module.prepare(source, self.folder / 'prepared')
        row = self.output_rows(result)[0]
        self.assertNotIn('vp_voltage_v', row)
        self.assertEqual(result['tables'][0]['ambiguous'][0]['field'], 'voltage')
        self.assertFalse(result['tables'][0]['numeric_ready'])

    def test_unknown_units_do_not_follow_magnitudes(self):
        source = self.make('Record,Time,Voltage,Current(mA/cm2)\n1,60,4200,1\n')
        result = module.prepare(source, self.folder / 'prepared')
        row = self.output_rows(result)[0]
        self.assertEqual(row['vp_voltage_unresolved'], '4200')
        self.assertEqual(row['vp_time_unresolved'], '60')
        self.assertEqual(row['vp_current_unresolved'], '1')
        self.assertFalse(result['tables'][0]['numeric_ready'])

    def test_unit_conflicting_with_identifier_is_not_erased(self):
        source = self.make('Cycle(Ah),Capacity(mAh)\n1,2\n')
        result = module.prepare(source, self.folder / 'prepared')
        self.assertEqual(self.output_rows(result)[0]['vp_cycle_unresolved'], '1')
        self.assertFalse(result['tables'][0]['numeric_ready'])

    def test_fractional_negative_and_nonfinite_cycles_are_not_valid_counts(self):
        source = self.make('Cycle,Capacity(mAh)\n1.0,2\n1.5,2\n-1,2\nNaN,2\n')
        result = module.prepare(source, self.folder / 'prepared')
        rows = self.output_rows(result)
        self.assertEqual(rows[0]['Cycle'], '1.0')
        self.assertEqual(rows[0]['vp_cycle'], '1')
        self.assertEqual([r['vp_cycle'] for r in rows[1:]], ['', '', ''])
        self.assertEqual(len(result['tables'][0]['conversion_issues']), 3)

    def test_missing_zero_nan_and_invalid_values_are_distinct(self):
        source = self.make('Record,Time(s),Voltage(V)\n1,0,0\n2,1,\n3,2,NaN\n4,3,bad\n5,4,Infinity\n')
        result = module.prepare(source, self.folder / 'prepared')
        rows = self.output_rows(result)
        self.assertEqual(rows[0]['vp_voltage_v'], '0')
        self.assertEqual([r['Voltage(V)'] for r in rows[1:]], ['', 'NaN', 'bad', 'Infinity'])
        self.assertEqual(len(rows), 5)
        self.assertEqual([v['kind'] for v in result['tables'][0]['missing_values']], ['empty', 'source_missing_token'])
        self.assertEqual(len(result['tables'][0]['conversion_issues']), 3)
        self.assertFalse(result['tables'][0]['numeric_ready'])

    def test_ragged_rows_retained_for_inspection_and_preparation_refused(self):
        source = self.make('Record,Time(s),Voltage(V)\n1,0,3.5\n2,1\n3,2,3.6,extra\n')
        result = module.inspect(source)
        self.assertEqual(len(result['tables'][0]['rows']), 3)
        self.assertEqual(result['tables'][0]['rows'][2][-1], 'extra')
        self.assertEqual(len(result['tables'][0]['issues']), 2)
        with self.assertRaisesRegex(ValueError, 'Ragged'):
            module.prepare(source, self.folder / 'prepared')
        self.assertFalse((self.folder / 'prepared').exists())

    def test_malformed_quotes_raise_before_any_output(self):
        source = self.make('Record,Voltage(V)\n1,"unfinished\n')
        with self.assertRaisesRegex(ValueError, 'Malformed'):
            module.prepare(source, self.folder / 'prepared')
        self.assertFalse((self.folder / 'prepared').exists())

    def test_duplicate_column_labels_require_an_explicit_export_fix(self):
        source = self.make('Cycle,Capacity(mAh),Capacity(mAh)\n1,2,3\n')
        with self.assertRaisesRegex(ValueError, 'No unambiguous'):
            module.prepare(source, self.folder / 'prepared')
        self.assertFalse((self.folder / 'prepared').exists())

    def test_protocol_is_not_a_measurement_table(self):
        source = self.make('StepNo,ObjectType,MainPara,LimitPara,EndCond,LogCond,ProtectCond\n1,Rest,,,t ~ 01:00:00,dT=00:30,\n')
        result = module.prepare(source, self.folder / 'prepared')
        self.assertEqual(result['status'], 'non_measurement_input')
        self.assertEqual(result['tables'][0]['role'], 'protocol')
        self.assertFalse((self.folder / 'prepared').exists())

    def test_step_and_cycle_headers_are_classified_without_merging(self):
        source = self.make('Step Index,Status,Time(s),Capacity(mAh)\n1,Rest,60,0\nCycle,Charge_Capacity(mAh),Discharge_Capacity(mAh)\n1,1,0.99\n')
        result = module.prepare(source, self.folder / 'prepared')
        self.assertEqual([t['role'] for t in result['tables']], ['step', 'cycle'])
        self.assertEqual(len(list((self.folder / 'prepared').glob('table-*.csv'))), 2)

    def test_source_provenance_collision_is_rejected(self):
        source = self.make('Record,Time(s),Voltage(V),vp_source_row\n1,0,3.5,99\n')
        with self.assertRaisesRegex(ValueError, 'collides'):
            module.prepare(source, self.folder / 'prepared')
        self.assertFalse((self.folder / 'prepared').exists())

    def test_canonical_collision_is_rejected(self):
        source = self.make('Cycle,Capacity(mAh/g),Specific Capacity(mAh/g)\n1,150,150\n')
        with self.assertRaisesRegex(ValueError, 'collapse'):
            module.prepare(source, self.folder / 'prepared')
        self.assertFalse((self.folder / 'prepared').exists())

    def test_output_directory_is_never_overwritten(self):
        source = self.make('Cycle,Capacity(mAh)\n1,2\n')
        target = self.folder / 'prepared'
        target.mkdir()
        sentinel = target / 'author.txt'
        sentinel.write_text('keep', encoding='utf-8')
        with self.assertRaises(FileExistsError):
            module.prepare(source, target)
        self.assertEqual(sentinel.read_text(encoding='utf-8'), 'keep')

    def test_native_cex_and_nda_return_export_action_without_fake_csv(self):
        for suffix in ('.cex', '.nda'):
            with self.subTest(suffix=suffix):
                source = self.folder / ('source' + suffix)
                source.write_bytes(b'not a tested native measurement')
                result = module.prepare(source, self.folder / ('out' + suffix))
                self.assertEqual(result['status'], 'needs_vendor_export')
                self.assertFalse(result['native_decoder'])
                self.assertFalse((self.folder / ('out' + suffix)).exists())

    def test_binary_content_with_csv_suffix_is_not_treated_as_export(self):
        source = self.folder / 'input.csv'
        source.write_bytes(b'LAND\x00binary')
        with self.assertRaisesRegex(ValueError, 'Binary'):
            module.inspect(source)

    def test_invalid_duration_keeps_raw_value(self):
        source = self.make('Record,Time(h:min:s),Voltage(V)\n1,01:60:00,3.5\n')
        result = module.prepare(source, self.folder / 'prepared')
        row = self.output_rows(result)[0]
        self.assertEqual(row['Time(h:min:s)'], '01:60:00')
        self.assertEqual(row['vp_time_s'], '')
        self.assertFalse(result['tables'][0]['numeric_ready'])

    @unittest.skipUnless(has('openpyxl'), 'XLSX dependency not installed')
    def test_xlsx_all_sheets_metadata_and_formula_preservation(self):
        from openpyxl import Workbook
        book = Workbook()
        info = book.active; info.title = 'Info'; info.append(['Instrument', 'Synthetic'])
        cycle = book.create_sheet('Cycle'); cycle.append(['Channel: synthetic']); cycle.append(['Cycle', 'Capacity(mAh)']); cycle.append([1, 2])
        record = book.create_sheet('Record'); record.append(['Index', 'Time(s)', 'Voltage(V)', 'Current(mA)']); record.append([1, 0, '=1+2', -1])
        schedule = book.create_sheet('Schedule'); schedule.append(['StepNo', 'ObjectType', 'MainPara', 'EndCond']); schedule.append([1, 'Rest', None, 't=60'])
        source = self.folder / 'source.xlsx'; book.save(source)
        result = module.prepare(source, self.folder / 'prepared')
        self.assertEqual(len(result['sheets']), 4)
        self.assertEqual([t['role'] for t in result['tables']], ['unclassified', 'cycle', 'record', 'protocol'])
        self.assertEqual(result['tables'][1]['source_rows'], [3])
        self.assertFalse(result['tables'][2]['numeric_ready'])
        with (self.folder / 'prepared' / result['tables'][2]['output']).open(encoding='utf-8', newline='') as stream:
            row = next(csv.DictReader(stream))
        self.assertEqual(row['Voltage(V)'], '=1+2')
        self.assertEqual(row['vp_voltage_v'], '')
        self.assertEqual((self.folder / 'prepared/source-original.xlsx').read_bytes(), source.read_bytes())

    @unittest.skipUnless(has('xlrd') and has('xlwt'), 'Legacy XLS test writer/reader unavailable')
    def test_legacy_xls_tables_are_not_renamed(self):
        import xlwt
        book = xlwt.Workbook()
        for name, headers, values in [('Cycle', ['Cycle', 'Capacity(Ah)'], [1, .002]),
                                     ('Record', ['Index', 'Time(s)', 'Voltage(V)'], [1, 0, 3.5])]:
            sheet = book.add_sheet(name)
            for col, value in enumerate(headers): sheet.write(0, col, value)
            for col, value in enumerate(values): sheet.write(1, col, value)
        source = self.folder / 'legacy.xls'; book.save(str(source))
        result = module.prepare(source, self.folder / 'prepared')
        self.assertEqual([t['role'] for t in result['tables']], ['cycle', 'record'])
        self.assertEqual(self.output_rows(result)[0]['vp_capacity_mah'], '2.000')
        self.assertEqual((self.folder / 'prepared/source-original.xls').read_bytes(), source.read_bytes())

    def test_cli_errors_are_concise_without_traceback(self):
        source = self.make('Record,Time(s),Voltage(V)\n1,0\n')
        result = subprocess.run([sys.executable, str(SCRIPT), 'prepare', '--input', str(source), '--output-dir', str(self.folder / 'prepared')],
                                capture_output=True, text=True, encoding='utf-8', env={**os.environ, 'PYTHONUTF8': '1'})
        self.assertEqual(result.returncode, 2)
        self.assertNotIn('Traceback', result.stderr)
        self.assertEqual(json.loads(result.stderr)['status'], 'error')

    @unittest.skipUnless(has('openpyxl') and has('xlrd'), 'Excel reader dependencies unavailable')
    def test_malformed_workbooks_report_errors_without_traceback(self):
        for suffix in ('.xlsx', '.xls'):
            with self.subTest(suffix=suffix):
                source = self.folder / ('broken' + suffix)
                source.write_bytes(b'broken workbook')
                result = subprocess.run([sys.executable, str(SCRIPT), 'inspect', '--input', str(source)],
                    capture_output=True, text=True, encoding='utf-8', env={**os.environ, 'PYTHONUTF8': '1'})
                self.assertEqual(result.returncode, 2)
                self.assertNotIn('Traceback', result.stderr)
                self.assertEqual(json.loads(result.stderr)['status'], 'error')


class NativeSourceLayerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.path = self.folder / 'synthetic.ndax'

    def tearDown(self):
        self.temp.cleanup()

    def test_absent_reader_returns_specific_local_action(self):
        self.path.write_bytes(b'PK')
        with patch.object(importlib.metadata, 'version', side_effect=importlib.metadata.PackageNotFoundError):
            result = module.prepare(self.path, self.folder / 'prepared')
        self.assertEqual(result['reason_code'], 'native_dependency_missing')
        self.assertFalse((self.folder / 'prepared').exists())

    def test_unchecked_reader_version_is_not_used(self):
        self.path.write_bytes(b'PK')
        with patch.object(importlib.metadata, 'version', return_value='9999.1'):
            result = module.inspect(self.path)
        self.assertEqual(result['reason_code'], 'native_dependency_version')

    @unittest.skipUnless(has_pinned_reader(), 'Optional pinned native reader not installed')
    def test_malformed_native_archive_returns_concise_fallback(self):
        self.path.write_bytes(b'PKbroken')
        result = module.prepare(self.path, self.folder / 'prepared')
        self.assertEqual(result['reason_code'], 'malformed_native_archive')
        self.assertFalse((self.folder / 'prepared').exists())

    @unittest.skipUnless(has_pinned_reader(), 'Optional pinned native reader not installed')
    def test_native_layers_preserve_zero_and_conflicting_timestamp(self):
        synthetic_ndax(self.path)
        before = self.path.read_bytes()
        result = module.prepare(self.path, self.folder / 'prepared')
        info = result['native_validation']
        self.assertEqual(result['status'], 'native_table_export')
        self.assertEqual(info['source_layer_counts'], {'measurements': 3, 'run_information': 4, 'native_steps': 2})
        self.assertEqual(info['reader_interpreted_records'], 2)
        self.assertEqual(info['zero_voltage_records_retained'], 1)
        self.assertEqual(info['duplicate_run_information_rows_retained'], 1)
        self.assertEqual(info['conflicting_run_information_ids'], 1)
        self.assertEqual(info['unit_and_timestamp_comparison'], 'pass')
        self.assertFalse(info['reader_arguments']['software_cycle_number'])
        self.assertFalse(info['interpolation'])
        self.assertFalse(info['merged_records'])
        self.assertFalse(result['ready_for_plot'])
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual((self.folder / 'prepared/source-original.ndax').read_bytes(), before)
        with (self.folder / 'prepared' / result['tables'][1]['output']).open(encoding='utf-8', newline='') as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual([row['Time(ms)'] for row in rows[:2]], ['0', '1000'])
        self.assertEqual([row['Index'] for row in rows[:2]], ['1', '1'])
        self.assertEqual(rows[1]['vp_charge_capacity_mah'], str(module.Decimal(rows[1]['Charge_Capacity(Ah)']) * 1000))
        with (self.folder / 'prepared' / result['tables'][2]['output']).open(encoding='utf-8', newline='') as stream:
            steps = list(csv.DictReader(stream))
        self.assertEqual(steps[0]['Cycle_ZeroBased'], '0')
        self.assertEqual(steps[0]['vp_cycle_zero_based'], '0')

    @unittest.skipUnless(has_pinned_reader(), 'Optional pinned native reader not installed')
    def test_native_missing_run_information_refuses_interpolation(self):
        synthetic_ndax(self.path, zero_voltage=False, missing_run=True)
        result = module.prepare(self.path, self.folder / 'prepared')
        self.assertEqual(result['status'], 'needs_vendor_export')
        self.assertEqual(result['reason_code'], 'native_interpolation_required')
        self.assertFalse((self.folder / 'prepared').exists())

    @unittest.skipUnless(has_pinned_reader(), 'Optional pinned native reader not installed')
    def test_internal_unreferenced_zero_is_not_called_padding(self):
        synthetic_ndax(self.path, unreferenced_zero=True)
        result = module.prepare(self.path, self.folder / 'prepared')
        self.assertEqual(result['reason_code'], 'native_unreferenced_zero_slot')
        self.assertFalse((self.folder / 'prepared').exists())

    @unittest.skipUnless(has_pinned_reader(), 'Optional pinned native reader not installed')
    def test_quantity_crosscheck_rejects_a_scale_mismatch(self):
        import NewareNDA
        synthetic_ndax(self.path)
        real_read = NewareNDA.read
        def mismatched(*args, **kwargs):
            frame = real_read(*args, **kwargs)
            frame['Voltage'] = frame['Voltage'] * 1000
            return frame
        with patch.object(NewareNDA, 'read', side_effect=mismatched):
            result = module.prepare(self.path, self.folder / 'prepared')
        self.assertEqual(result['reason_code'], 'native_decoder_quantity_mismatch')
        self.assertFalse((self.folder / 'prepared').exists())

    @unittest.skipUnless(has_pinned_reader(), 'Optional pinned native reader not installed')
    def test_unknown_native_version_requires_vendor_export(self):
        synthetic_ndax(self.path, version=17)
        result = module.inspect(self.path)
        self.assertEqual(result['reason_code'], 'unvalidated_native_version')

    @unittest.skipUnless(has_pinned_reader(), 'Optional pinned native reader not installed')
    def test_auxiliary_channels_do_not_disappear_silently(self):
        synthetic_ndax(self.path, auxiliary=True)
        result = module.inspect(self.path)
        self.assertEqual(result['reason_code'], 'unvalidated_native_auxiliary')


if __name__ == '__main__':
    unittest.main()
