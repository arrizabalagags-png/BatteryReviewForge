"""Inspect uploaded tables before assigning scientific meaning to their columns."""

from __future__ import annotations

import csv
from pathlib import Path

from .data import DataContractError


PLOT_COLUMNS = {
    "coulombic_efficiency": (("series", "cycle", "ce_pct"), ("series", "cycle", "ce_numerator", "ce_denominator")),
    "full_cell_cycling": (("series", "cycle", "discharge_capacity"),),
    "half_cell_cycling": (("series", "cycle", "discharge_capacity"),),
    "symmetric_cell_voltage": (("series", "time_h", "voltage_mv"),),
    "voltage_capacity": (("series", "cycle", "direction", "capacity", "voltage_v"),),
    "nyquist": (("series", "z_real_ohm", "minus_z_imag_ohm"),),
    "cycle_retention": (("series", "cycle", "retention_pct"),),
    "rate_capability": (("series", "step", "rate_label", "capacity"),),
    "tofsims_map": (("sample_id", "fragment", "x_um", "y_um", "signal"),),
    "tofsims_depth": (("sample_id", "fragment", "sputter_time_s", "signal"),),
}

ALIASES = {
    "cycle": ("cycle number", "cycle index", "循环圈数", "圈数"),
    "ce_pct": ("coulombic efficiency (%)", "ce (%)", "库伦效率"),
    "discharge_capacity": ("discharge capacity", "放电容量", "放电比容量"),
    "time_h": ("time (h)", "elapsed time (h)", "时间/h"),
    "voltage_mv": ("voltage (mv)", "电压/mv"),
    "voltage_v": ("voltage (v)", "电压/v"),
    "z_real_ohm": ("zreal", "z' (ohm)", "re(z)"),
    "minus_z_imag_ohm": ("-zimag", "-z'' (ohm)", "-im(z)"),
}


def _check_rows(rows: list[dict]) -> list[dict]:
    if not rows:
        raise DataContractError("The selected table has no data rows")
    headers = list(rows[0])
    if not headers or len(headers) != len(set(headers)) or any(not str(h).strip() for h in headers):
        raise DataContractError("Table needs nonempty, unique column headers")
    if any(None in row for row in rows):
        raise DataContractError("A data row has more cells than the header; fix the source table")
    return rows


def read_table(path: str | Path, *, sheet: str | None = None) -> list[dict]:
    """Read a flat CSV/TSV/TXT/XLSX sheet; never guess a header offset or cell meaning."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv", ".txt"}:
        content = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                content = path.read_text(encoding=encoding)
                break
            except UnicodeDecodeError:
                pass
        if content is None:
            raise DataContractError("Cannot decode text table; export UTF-8 CSV")
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=",\t;")
        except csv.Error:
            dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
        reader = csv.DictReader(content.splitlines(), dialect=dialect)
        if not reader.fieldnames or None in reader.fieldnames or len(reader.fieldnames) != len(set(reader.fieldnames)):
            raise DataContractError("Text table needs unique headers on its first row")
        return _check_rows(list(reader))
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise DataContractError("XLSX needs openpyxl; install the plotting requirements") from exc
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet is None and len(book.sheetnames) != 1:
                raise DataContractError(f"Workbook has multiple sheets {book.sheetnames}; choose --sheet")
            if sheet is not None and sheet not in book.sheetnames:
                raise DataContractError(f"Sheet {sheet!r} not found; available: {book.sheetnames}")
            worksheet = book[sheet or book.sheetnames[0]]
            iterator = worksheet.iter_rows(values_only=True)
            header = next(iterator, None)
            if not header:
                raise DataContractError("Selected sheet has no header")
            headers = [str(item).strip() if item is not None else "" for item in header]
            if not all(headers) or len(headers) != len(set(headers)):
                raise DataContractError("XLSX needs unique headers on its first row")
            rows = []
            for line in iterator:
                if any(item is not None and str(item).strip() for item in line):
                    rows.append({name: "" if index >= len(line) or line[index] is None else str(line[index])
                                 for index, name in enumerate(headers)})
            return _check_rows(rows)
        finally:
            book.close()
    raise DataContractError(f"Unsupported table {suffix or '(no extension)'}; export CSV/TSV/TXT/XLSX")


def inspect_table(rows: list[dict]) -> dict:
    """Return observable columns and chart candidates, never invent units or test conditions."""
    _check_rows(rows)
    headers = list(rows[0])
    present = set(headers)
    candidates = {
        kind: [list(option) for option in options if set(option) <= present]
        for kind, options in PLOT_COLUMNS.items()
    }
    suggestions = {
        canonical: [header for header in headers if header.strip().lower() in names]
        for canonical, names in ALIASES.items()
    }
    return {
        "row_count": len(rows), "columns": headers,
        "preview": rows[:3],
        "matching_plot_types": {key: value for key, value in candidates.items() if value},
        "column_mapping_suggestions": {key: value for key, value in suggestions.items() if value},
        "note": "Candidates are column-shape matches, not scientific validation. Supply source IDs, units, cell/test conditions and evidence state.",
    }


def apply_mapping(rows: list[dict], columns: dict[str, str], common: dict) -> list[dict]:
    if not isinstance(columns, dict) or not isinstance(common, dict):
        raise DataContractError("Metadata needs columns and common objects")
    available = set(rows[0])
    missing = {source for source in columns.values() if source not in available}
    if missing:
        raise DataContractError(f"Column mapping refers to missing headers: {sorted(missing)}")
    mapped = []
    for row in rows:
        record = {**common, **row}
        for canonical, source in columns.items():
            record[canonical] = row[source]
        mapped.append(record)
    return mapped
