"""Import local exports while keeping original instrument observations.

Optional native import is restricted to the checked NDAX NDC-14 layout.
"""
from __future__ import annotations
import argparse
import csv
import hashlib
import importlib.util
import io
import json
import re
import shutil
import sys
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path

MAX_SOURCE_BYTES = 64 * 1024 * 1024
NATIVE = {
    '.cex': ('LANDdt', '在 LANDdt 中打开原文件，导出数据表或复制数据到 Excel，再保存为 XLSX、CSV 或制表符 TXT。'),
    '.nda': ('BTSDA', 'NDA 尚无本项目实样验证；在对应版本 BTSDA 中导出循环层、工步层和记录层的 Excel 或文本表。'),
    '.ndax': ('BTSDA', '在对应版本 BTSDA 中打开原文件，导出完整循环层、工步层和记录层的 Excel 或文本表。'),
}
ALIASES = {
    'cycle': ('cycle', 'cycle index', 'cycle number', '循环', '循环号', '循环序号', '循环圈数', '圈数'),
    'cycle_zero_based': ('cycle zero based',),
    'step': ('step', 'step number', 'stepno', '工步', '工步号', '工步序号', '过程', '过程号'),
    'step_index': ('step index',),
    'record': ('record', 'record id', 'record index', 'index', 'data point', '记录', '记录号', '记录序号', '数据点'),
    'step_type': ('step type', 'step name', 'status', '工步类型', '工步名称', '过程模式', '状态'),
    'time': ('time', 'test time', 'total time', 'relative time', 'step time', 'elapsed time', '时间', '测试时间', '总时间', '相对时间', '工步时间', '过程时间'),
    'voltage': ('voltage', 'voltage avg', 'voltage average', '电压', '平均电压'),
    'current': ('current', '电流'),
    'capacity': ('capacity', '容量'),
    'charge_capacity': ('charge capacity', 'chg capacity', '充电容量'),
    'discharge_capacity': ('discharge capacity', 'dchg capacity', '放电容量'),
    'specific_capacity': ('specific capacity', '比容量'),
    'charge_specific_capacity': ('charge specific capacity', '充电比容量'),
    'discharge_specific_capacity': ('discharge specific capacity', '放电比容量'),
    'ce': ('coulombic efficiency', 'coulombic efficiency ratio', 'ce', '库伦效率', '库仑效率', '充放电效率'),
    'energy': ('energy', '能量'),
    'charge_energy': ('charge energy', '充电能量'),
    'discharge_energy': ('discharge energy', '放电能量'),
}
COUNT_FIELDS = {'cycle', 'cycle_zero_based', 'step', 'step_index', 'record'}
CONTROL_HEADERS = {'mainpara', 'limitpara', 'endcond', 'logcond', 'protectcond',
                   'controlparameter', 'endcondition', 'logcondition', 'protection',
                   '主参数', '限制参数', '截止条件', '记录条件', '保护条件'}
MISSING_TOKENS = {'nan', 'na', 'n/a', 'null', 'none', '--'}
PENDING_CONTEXT = [
    '确认实际电池配置、通道及续接测试身份。',
    '按目标图确认循环/工步定义、分支含义及时间基准；不由电流正负自动指定。',
    '确认容量/能量的计量基准和归一化分母；质量、面积未知时不换算。',
    '按比较范围确认温度、电压窗口、倍率及其他实验条件。',
    '重算库伦效率前确认对应循环和协议分母。',
]


def norm(text):
    return re.sub(r'[\s_]+', '', unicodedata.normalize('NFKC', str(text)).casefold()).replace('μ', 'u').replace('µ', 'u').replace('−', '-')


def identify(header):
    """Recognise a label; missing or unsupported units stay unresolved."""
    text = unicodedata.normalize('NFKC', str(header)).strip()
    match = re.fullmatch(r'(.*?)\s*[\[(]([^\])]+)[\])]\s*', text)
    if match:
        base, unit = match.groups()
    elif '/' in text:
        base, unit = text.split('/', 1)
    else:
        base, unit = text, ''
    for canonical, aliases in ALIASES.items():
        if norm(base) in {norm(a) for a in aliases}:
            return canonical, unit.strip()
    return None, unit.strip()


def cell(value):
    if value is None:
        return ''
    if hasattr(value, 'isoformat'):
        return value.isoformat()
    return str(value)


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def matrices(path):
    """Read logical cells: no formula evaluation, merged fill or row deletion."""
    path = Path(path)
    if path.stat().st_size > MAX_SOURCE_BYTES:
        raise ValueError('File exceeds 64 MiB; export separate channels or recording ranges.')
    suffix = path.suffix.lower()
    if suffix in NATIVE:
        raise ValueError(NATIVE[suffix][1])
    if suffix in {'.csv', '.tsv', '.txt'}:
        data = path.read_bytes()
        if b'\x00' in data[:8192] and not data.startswith((b'\xff\xfe', b'\xfe\xff')):
            raise ValueError('Binary content is not a text table; use the instrument export function.')
        if data.startswith((b'\xff\xfe', b'\xfe\xff')):
            encoding = 'utf-16'
            content = data.decode(encoding)
        else:
            for encoding in ('utf-8-sig', 'gb18030'):
                try:
                    content = data.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                raise ValueError('Text encoding unknown; export UTF-8 CSV.')
        sample = content[:16384]
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=',\t;').delimiter
        except csv.Error:
            delimiter = max(('\t', ',', ';'), key=lambda value: sample.count(value))
        try:
            rows = list(csv.reader(io.StringIO(content, newline=''), delimiter=delimiter, strict=True))
        except csv.Error as error:
            raise ValueError('Malformed quoted text table; export a complete CSV or TXT table.') from error
        return [('text', rows, {'format': 'text', 'encoding': encoding, 'delimiter': delimiter})]
    if suffix == '.xlsx':
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise ValueError('Install requirements.txt in the project environment to read XLSX.') from error
        try:
            book = load_workbook(path, read_only=True, data_only=False)
        except Exception as error:
            raise ValueError('XLSX cannot be read as a complete workbook; retain the original and export again.') from error
        try:
            return [(ws.title, [[cell(v) for v in line] for line in ws.iter_rows(values_only=True)],
                     {'format': 'XLSX', 'formula_evaluation': False}) for ws in book.worksheets]
        except Exception as error:
            raise ValueError('XLSX sheet data cannot be read completely; retain the original and export again.') from error
        finally:
            book.close()
    if suffix == '.xls':
        try:
            import xlrd
        except ImportError as error:
            raise ValueError('Install xlrd>=2,<3 in the project environment for legacy XLS; do not rename it.') from error
        try:
            book = xlrd.open_workbook(str(path))
        except Exception as error:
            raise ValueError('XLS cannot be read as a legacy workbook; retain the original and export again.') from error
        result = []
        try:
            for ws in book.sheets():
                rows = []
                for r in range(ws.nrows):
                    line = []
                    for c in range(ws.ncols):
                        value = ws.cell_value(r, c)
                        if ws.cell_type(r, c) == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, book.datemode).isoformat()
                        line.append(cell(value))
                    rows.append(line)
                result.append((ws.name, rows, {'format': 'XLS', 'date_mode': book.datemode,
                    'formula_values': 'stored results only; not recalculated'}))
            return result
        except Exception as error:
            raise ValueError('XLS sheet data cannot be read completely; retain the original and export again.') from error
        finally:
            book.release_resources()
    raise ValueError('Supported exports: CSV, TSV, TXT, XLSX and XLS; checked native NDAX is optional.')


def trimmed(row):
    row = list(map(cell, row))
    while row and not row[-1].strip():
        row.pop()
    return row


def valid_header(row):
    if len(row) < 2 or any(not str(v).strip() for v in row) or len(set(map(str, row))) != len(row):
        return False
    for value in row:
        try:
            Decimal(str(value))
        except InvalidOperation:
            continue
        return False
    return True


def is_protocol(headers):
    names = {norm(h) for h in headers}
    return len(names & CONTROL_HEADERS) >= 2 or {'objecttype', 'mainpara'} <= names


def table_blocks(sheet, rows, format_info):
    workbook = format_info.get('format') in {'XLSX', 'XLS'}
    header_rows = [trimmed(row) if workbook else list(map(cell, row)) for row in rows]
    candidates = [i for i, row in enumerate(header_rows) if valid_header(row) and
                  (is_protocol(row) or sum(identify(h)[0] is not None for h in row) >= 2)]
    if not candidates:
        first = next((i for i, row in enumerate(header_rows) if any(v.strip() for v in row)), None)
        if first is not None and valid_header(header_rows[first]):
            candidates = [first]
    tables = []
    for position, start in enumerate(candidates):
        end = candidates[position + 1] if position + 1 < len(candidates) else len(rows)
        headers = header_rows[start]
        body, source_rows, issues, blank_rows = [], [], [], []
        for i in range(start + 1, end):
            row = list(map(cell, rows[i]))
            if not any(v.strip() for v in row):
                blank_rows.append(i + 1)
                continue
            if workbook and len(row) > len(headers) and not any(v.strip() for v in row[len(headers):]):
                row = row[:len(headers)]
            if len(row) != len(headers):
                issues.append({'row': i + 1, 'reason': 'row width differs from the header; all cells retained for inspection'})
            body.append(row)
            source_rows.append(i + 1)
        tables.append(dict(sheet=sheet, header_row=start + 1, headers=headers, rows=body,
                           source_rows=source_rows, blank_source_rows=blank_rows, issues=issues, format=format_info))
    return tables


def declared_mapping(headers):
    result, ambiguous = {}, []
    for index, header in enumerate(headers):
        name, unit = identify(header)
        if name:
            result.setdefault(name, []).append(dict(index=index, header=header, unit=unit))
    for name, items in list(result.items()):
        if len(items) > 1:
            ambiguous.append({'field': name, 'headers': [item['header'] for item in items]})
            del result[name]
        else:
            result[name] = items[0]
    return result, ambiguous


def converted(name, unit, value):
    """Convert only declared units, retaining raw cells in separate columns."""
    value = cell(value)
    u = norm(unit).replace('·', '').replace('^', '')
    if name == 'step_type':
        return name, value, 'category'
    if name in COUNT_FIELDS:
        if u not in {'', 'count', 'counts', 'index', '1', '次', '圈'}:
            return name + '_unresolved', value, unit
        if not value.strip():
            return name, '', 'count'
        try:
            number = Decimal(value)
        except InvalidOperation as error:
            raise ValueError('identifier is not a decimal integer') from error
        if not number.is_finite() or number < 0 or number != number.to_integral_value():
            raise ValueError('identifier must be a finite nonnegative integer; cycle labels are not renumbered')
        return name, str(int(number)), 'count'
    target, factor = None, None
    if name == 'time':
        target = 'time_s'
        if u in {'h:min:s', 'h:m:s', 'hh:mm:ss', 'hh:mm:ss.sss'}:
            if not value.strip():
                return target, '', 's'
            match = re.fullmatch(r'(\d+):(\d{1,2}):(\d{1,2}(?:\.\d+)?)', value.strip())
            if not match or int(match[2]) >= 60 or Decimal(match[3]) >= 60:
                raise ValueError('duration must be nonnegative HH:MM:SS')
            return target, str(Decimal(match[1]) * 3600 + Decimal(match[2]) * 60 + Decimal(match[3])), 's'
        factor = {'s': 1, 'sec': 1, '秒': 1, 'min': 60, '分钟': 60, 'h': 3600, '小时': 3600, 'ms': Decimal('.001')}.get(u)
    elif name == 'voltage':
        target = 'voltage_v'; factor = {'v': 1, 'mv': Decimal('.001'), 'uv': Decimal('.000001')}.get(u)
    elif name == 'current':
        target = 'current_a'; factor = {'a': 1, 'ma': Decimal('.001'), 'ua': Decimal('.000001')}.get(u)
    elif name in {'capacity', 'charge_capacity', 'discharge_capacity'}:
        if u in {'mah/g', 'mahg-1', 'ah/kg', 'ahkg-1'}:
            target = name + '_mah_g'; factor = 1
        elif u in {'ah/g', 'ahg-1'}:
            target = name + '_mah_g'; factor = 1000
        else:
            target = name + '_mah'; factor = {'mah': 1, 'ah': 1000, 'uah': Decimal('.001')}.get(u)
    elif 'specific_capacity' in name:
        target = name.replace('specific_capacity', 'capacity') + '_mah_g'
        factor = {'mah/g': 1, 'mahg-1': 1, 'ah/kg': 1, 'ahkg-1': 1, 'ah/g': 1000, 'ahg-1': 1000}.get(u)
    elif name == 'ce':
        target = 'ce_pct'; factor = {'%': 1, 'percent': 1, 'fraction': 100, '1': 100}.get(u)
    elif name in {'energy', 'charge_energy', 'discharge_energy'}:
        target = name + '_wh'; factor = {'wh': 1, 'mwh': Decimal('.001')}.get(u)
    if factor is None:
        return name + '_unresolved', value, unit or 'unknown'
    output_unit = ('s' if name == 'time' else 'V' if name == 'voltage' else 'A' if name == 'current'
                   else '%' if name == 'ce' else 'Wh' if 'energy' in name else 'mAh g-1' if target.endswith('_mah_g') else 'mAh')
    if not value.strip():
        return target, '', output_unit
    if value.strip().casefold() in MISSING_TOKENS:
        raise ValueError('missing-value token retained in source; it is not zero')
    try:
        number = Decimal(value)
    except InvalidOperation as error:
        raise ValueError('not an unambiguous decimal number; formulas are not evaluated') from error
    if not number.is_finite():
        raise ValueError('nonfinite value retained in source; it is not filtered')
    if name == 'time' and number < 0:
        raise ValueError('negative elapsed duration requires source review')
    return target, str(number * Decimal(factor)), output_unit


def role_of(table):
    if table['format'].get('table_role'):
        return table['format']['table_role']
    if is_protocol(table['headers']):
        return 'protocol'
    names = set(table['mapping'])
    capacity_fields = {'capacity', 'discharge_capacity', 'charge_capacity', 'ce', 'specific_capacity', 'discharge_specific_capacity', 'charge_specific_capacity'}
    if 'cycle' in names and names & capacity_fields and not names & {'record', 'step', 'step_index'}:
        return 'cycle'
    if names & {'step', 'step_index'} and 'step_type' in names and 'record' not in names:
        return 'step'
    if 'record' in names or {'voltage', 'time'} <= names:
        return 'record'
    return 'unclassified'


def inspect(path):
    path = Path(path)
    if path.stat().st_size > MAX_SOURCE_BYTES:
        raise ValueError('File exceeds 64 MiB; export separate channels or recording ranges.')
    record = {'schema_version': 2, 'source': str(path.resolve()), 'sha256': sha256(path),
              'native_decoder': False, 'status': 'table_export', 'tables': [], 'sheets': [],
              'scientific_context_confirmed': False, 'ready_for_plot': False,
              'pending_context': list(PENDING_CONTEXT)}
    suffix = path.suffix.lower()
    if suffix in NATIVE:
        record.update(status='needs_vendor_export', software=NATIVE[suffix][0], next_step=NATIVE[suffix][1])
        if suffix != '.ndax':
            record['reason_code'] = 'unvalidated_native_format'
            return record
        spec = importlib.util.spec_from_file_location('vp_neware_ndax14', Path(__file__).with_name('neware_ndax14.py'))
        reader = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(reader)
        try:
            source_matrices, native_info = reader.read_source_layers(path)
        except reader.NativeExportNeeded as error:
            record.update(reason_code=error.code, reason=str(error), native_validation=error.details)
            return record
        record.update(status='native_table_export', native_decoder=True, native_validation=native_info)
        record['pending_context'].extend(native_info['pending_context'])
    else:
        source_matrices = matrices(path)
    for sheet, rows, info in source_matrices:
        tables = table_blocks(sheet, rows, info)
        record['sheets'].append({'name': sheet, 'source_rows': len(rows), 'tables': len(tables)})
        for table in tables:
            mapping, ambiguous = declared_mapping(table['headers'])
            table.update(mapping=mapping, ambiguous=ambiguous,
                         unmapped_source_headers=[h for h in table['headers'] if identify(h)[0] is None])
            table['role'] = role_of(table)
            table['scientific_context_confirmed'] = False
            record['tables'].append(table)
    if not record['tables']:
        raise ValueError('No unambiguous tabular header found. Keep the original file and inspect its export settings.')
    if all(t['role'] in {'protocol', 'unclassified'} for t in record['tables']):
        record.update(status='non_measurement_input', reason_code='protocol_or_information_table',
                      next_step='该文件是工步/测试程序或信息表；请提供实际循环、工步统计或记录层测量数据。')
    return record


def prepare(path, output):
    path, output = Path(path), Path(output)
    if output.exists():
        raise FileExistsError('Choose a new output directory; existing data are not overwritten.')
    record = inspect(path)
    if record['status'] not in {'table_export', 'native_table_export'}:
        return record
    for table in record['tables']:
        if table['issues']:
            raise ValueError('Ragged table contains rows that cannot be retained safely; inspect the source export before preparing.')
    outputs = []
    for number, table in enumerate(record['tables'], 1):
        mapping, new_fields, issues, converted_rows, missing = table['mapping'], {}, [], [], []
        for row, source_row in zip(table['rows'], table['source_rows']):
            new = {}
            for name, item in mapping.items():
                raw_value = row[item['index']]
                key, _, unit = converted(name, item['unit'], '')
                try:
                    key, value, unit = converted(name, item['unit'], raw_value)
                except ValueError as error:
                    value = ''
                    issues.append({'source_row': source_row, 'header': item['header'], 'reason': str(error)})
                if key in new:
                    raise ValueError('Several columns collapse to one canonical field; resolve the mapping explicitly.')
                new[key] = value
                new_fields[key] = {'source_header': item['header'], 'source_unit': item['unit'] or 'unknown', 'unit': unit}
                if name != 'step_type' and not raw_value.strip():
                    missing.append({'field': key, 'source_row': source_row, 'kind': 'empty'})
                elif name != 'step_type' and raw_value.strip().casefold() in MISSING_TOKENS:
                    missing.append({'field': key, 'source_row': source_row, 'kind': 'source_missing_token'})
            converted_rows.append(new)
        derived = list(new_fields)
        all_headers = ['vp_source_sheet', 'vp_source_row', *table['headers'], *['vp_' + key for key in derived]]
        if len(all_headers) != len(set(all_headers)):
            raise ValueError('Source header collides with a provenance or canonical field; choose an explicit mapping.')
        stream = io.StringIO(newline='')
        writer = csv.writer(stream)
        writer.writerow(all_headers)
        for row, source_row, new in zip(table['rows'], table['source_rows'], converted_rows):
            writer.writerow([table['sheet'], source_row, *row, *[new.get(key, '') for key in derived]])
        outputs.append((f'table-{number:02d}-{table["role"]}.csv', stream.getvalue()))
        numeric_ready = (bool(derived) and bool(table['rows']) and not issues and not missing and not table['ambiguous']
                         and not any(key.endswith('_unresolved') for key in derived)
                         and table['role'] not in {'protocol', 'unclassified'})
        table.update(normalised_fields=new_fields, conversion_issues=issues, missing_values=missing,
                     source_row_count=len(table['rows']), output=outputs[-1][0], numeric_ready=numeric_ready,
                     ready_for_plot=False)
        table.pop('rows')
    if record['sha256'] != sha256(path):
        raise ValueError('Source changed during inspection; rerun against a stable original file.')
    output.mkdir(parents=True, exist_ok=False)
    source_copy = output / ('source-original' + path.suffix.lower())
    shutil.copyfile(path, source_copy)
    if sha256(source_copy) != record['sha256']:
        raise ValueError('Source copy hash differs; preserve the original and inspect the incomplete output.')
    for name, content in outputs:
        (output / name).write_text(content, encoding='utf-8')
    record.update(source_copy=source_copy.name, original_file_unchanged=record['sha256'] == sha256(path),
                  scientific_validation=False,
                  transformations='Separate tables; source copied byte for byte; raw cells retained; vp_ fields use declared units. No smoothing, aggregation, branch/sign inference, cycle renumbering, mass/area normalisation or CE calculation.')
    (output / 'instrument-record.json').write_text(json.dumps(record, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return record


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('action', choices=('inspect', 'prepare'))
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--output-dir', type=Path)
    args = parser.parse_args()
    if args.action == 'prepare' and args.output_dir is None:
        parser.error('prepare needs --output-dir')
    try:
        report = inspect(args.input) if args.action == 'inspect' else prepare(args.input, args.output_dir)
    except (OSError, ValueError, ImportError) as error:
        print(json.dumps({'status': 'error', 'reason': str(error)}, ensure_ascii=False), file=sys.stderr)
        return 2
    if args.action == 'inspect':
        for table in report['tables']:
            rows = table.pop('rows')
            table['source_row_count'] = len(rows)
            table['preview'] = rows[:3]
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == '__main__':
    sys.exit(main())
